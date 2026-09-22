import base64
from io import BytesIO
from pathlib import Path
import zipfile

from PIL import Image

import main
from docx import Document
from fastapi.testclient import TestClient

from main import app, parse_html_to_sections


client = TestClient(app)


def test_health_and_config_do_not_disclose_key(test_data_dir: Path) -> None:
    health = client.get("/api/health")
    assert health.status_code == 200
    assert health.json()["status"] == "ok"
    assert Path(health.json()["database"]).is_relative_to(test_data_dir)
    config = client.get("/api/config").json()
    assert "gemini_api_key" not in config
    assert "api_key_masked" in config
    assert Path(config["database_path"]).is_relative_to(test_data_dir)


def test_html_parser_preserves_blocks_and_images() -> None:
    blocks = parse_html_to_sections(
        '<h1>Desarrollo</h1><p>Texto <strong>importante</strong>.</p>'
        '<ol><li>Paso uno</li><li>Paso dos</li></ol><img src="missing.png" alt="captura" />'
    )
    assert [block["type"] for block in blocks] == ["h1", "paragraph", "ordered_item", "ordered_item", "image"]
    assert "<strong>" in blocks[1]["content_html"]


def test_export_docx() -> None:
    response = client.post("/api/export-docx", json={
        "title": "Prueba de exportación",
        "subject": "Programación",
        "student": "Nicolás",
        "html_content": "<h1>Introducción</h1><p>Texto con <strong>formato</strong>.</p>",
    })
    assert response.status_code == 200
    assert response.content.startswith(b"PK")
    document = Document(BytesIO(response.content))
    paragraphs = document.paragraphs
    title = next(paragraph for paragraph in paragraphs if paragraph.text == "Prueba de exportación")
    assert title.style.name == "Title"
    introduction = next(paragraph for paragraph in paragraphs if paragraph.text == "Introducción")
    assert introduction.style.name == "Heading 1"
    formatted = next(paragraph for paragraph in paragraphs if paragraph.text == "Texto con formato.")
    assert any(run.text == "formato" and run.bold for run in formatted.runs)


def test_export_docx_with_table() -> None:
    response = client.post("/api/export-docx", json={
        "title": "Documento con Tabla",
        "subject": "Bases de Datos",
        "html_content": "<h2>Tabla Comparativa</h2><table><tr><th>Servicio</th><th>Puerto</th></tr><tr><td>VSFTPD</td><td>21</td></tr></table>",
    })
    assert response.status_code == 200
    assert response.content.startswith(b"PK")
    document = Document(BytesIO(response.content))
    assert len(document.tables) >= 2
    content_table = document.tables[-1]
    assert content_table.cell(0, 0).text == "Servicio"
    assert content_table.cell(1, 0).text == "VSFTPD"


def test_html_parser_preserves_images_nested_in_paragraphs() -> None:
    # Regression test for H-02: images inside <p> were silently dropped
    html = '<p>Nota previa: <img src="/api/evidence/test/content" alt="Captura Red Hat" /></p>'
    blocks = parse_html_to_sections(html)
    assert len(blocks) == 2
    assert blocks[0]["type"] == "paragraph"
    assert "Nota previa:" in blocks[0]["content"]
    assert blocks[1]["type"] == "image"
    assert blocks[1]["image_path"] == "/api/evidence/test/content"
    assert blocks[1]["image_alt"] == "Captura Red Hat"


def test_html_parser_handles_standalone_paragraph_image() -> None:
    html = '<p><img src="test.png" alt="Diagrama" /></p>'
    blocks = parse_html_to_sections(html)
    assert len(blocks) == 1
    assert blocks[0]["type"] == "image"
    assert blocks[0]["image_path"] == "test.png"


def test_html_parser_deduplicates_redundant_caption_paragraph() -> None:
    # Reproduces the exact HTML injected by EditorView.jsx on upload
    html = (
        '<h2>Desarrollo</h2>'
        '<img src="/api/evidence/test/content" alt="Ejecución de servicios" />'
        '<p><em>Figura: Ejecución de servicios</em></p>'
        '<p>Texto posterior del informe.</p>'
    )
    blocks = parse_html_to_sections(html)
    assert len(blocks) == 3
    assert blocks[0]["type"] == "h2"
    assert blocks[1]["type"] == "image"
    assert blocks[1]["image_alt"] == "Figura: Ejecución de servicios"
    assert blocks[2]["type"] == "paragraph"
    assert blocks[2]["content"] == "Texto posterior del informe."


def test_version_endpoints() -> None:
    from main import store
    project_id = store.create_project(
        title="Test Version API", subject="Sistemas", student="Nicolas", career="Info",
        summary="Test", rubric_text="Test", deliverables=[], outline=[]
    )
    # Create snapshots
    snap1 = client.post(f"/api/projects/{project_id}/snapshots", json={"html_content": "<p>V1</p>", "reason": "manual"})
    assert snap1.status_code == 200
    assert snap1.json()["version"] == 1

    snap2 = client.post(f"/api/projects/{project_id}/snapshots", json={"html_content": "<p>V2</p>", "reason": "manual"})
    assert snap2.status_code == 200
    assert snap2.json()["version"] == 2

    # List versions
    v_list = client.get(f"/api/projects/{project_id}/versions")
    assert v_list.status_code == 200
    assert len(v_list.json()) >= 2
    assert v_list.json()[0]["version_number"] == 2

    # Restore version 1
    restore_resp = client.post(
        f"/api/projects/{project_id}/versions/1/restore",
        json={"current_html_content": "<p>Borrador no guardado</p>"},
    )
    assert restore_resp.status_code == 200
    data = restore_resp.json()
    assert data["restored_from"] == 1
    assert "<p>V1</p>" in data["html_content"]
    backup = store.get_version(project_id, data["backup_version_number"])
    assert backup is not None
    assert backup["reason"] == "before_restore_v1"
    assert backup["html_content"] == "<p>Borrador no guardado</p>"


def test_quick_create_and_patch_section() -> None:
    # 1. Quick create without calling Gemini
    create_resp = client.post("/api/projects/quick-create", json={
        "title": "Proyecto Red Hat Offline",
        "subject": "Sistemas Operativos",
        "student": "Nicolás Jara",
        "pauta_text": "Instalar VSFTPD y MariaDB en Red Hat",
        "deliverables": ["Informe Word con capturas"],
        "criteria": [
            {
                "criterion": "Servicios",
                "indicator": "Configura VSFTPD y MariaDB",
                "points": 30.0,
                "requires_evidence": True,
            }
        ]
    })
    assert create_resp.status_code == 200
    pid = create_resp.json()["project_id"]
    assert pid.startswith("prj_")

    # 2. Patch a specific section
    patch_resp = client.patch(f"/api/projects/{pid}/section", json={
        "heading": "Desarrollo Técnico",
        "content_html": "<p>Se ejecutó el comando <code>systemctl start vsftpd</code> con éxito.</p>",
        "create_snapshot": True,
        "reason": "agent:add_vsftpd"
    })
    assert patch_resp.status_code == 200
    doc_html = patch_resp.json()["html_content"]
    assert "Desarrollo Técnico" in doc_html
    assert "systemctl start vsftpd" in doc_html

    # 3. Update existing section without overwriting other sections
    patch_resp2 = client.patch(f"/api/projects/{pid}/section", json={
        "heading": "Conclusión",
        "content_html": "<p>Se cumplieron todos los objetivos.</p>",
        "create_snapshot": True
    })
    assert patch_resp2.status_code == 200
    doc2_html = patch_resp2.json()["html_content"]
    assert "systemctl start vsftpd" in doc2_html
    assert "Conclusión" in doc2_html
    assert "Se cumplieron todos los objetivos." in doc2_html


def test_audit_only_reports_stored_evidence_for_the_exact_criterion(monkeypatch) -> None:
    create_resp = client.post("/api/projects/quick-create", json={
        "title": "Auditoría con evidencia verificable",
        "criteria": [{
            "criterion": "Servicios",
            "indicator": "Configura VSFTPD y MariaDB",
            "points": 30,
            "requires_evidence": True,
        }],
    })
    assert create_resp.status_code == 200
    project_id = create_resp.json()["project_id"]
    criterion_id = create_resp.json()["criteria"][0]["id"]
    html = "<p>Configura VSFTPD y MariaDB. La captura y la figura muestran el resultado.</p>"

    local_audit = client.post(
        f"/api/projects/{project_id}/audit",
        json={"html_content": html, "use_ai": False},
    )
    assert local_audit.status_code == 200
    local_item = local_audit.json()["items"][0]
    assert local_item["evidence_found"] is False
    assert local_item["status"] == "partial"

    monkeypatch.setattr(main, "get_gemini_api_key", lambda: "test-key")
    monkeypatch.setattr(main.ai, "audit", lambda prompt: main.AuditResult(
        summary="La IA afirma cumplimiento total.",
        estimated_score=100,
        items=[main.AuditItem(
            criterion_id=criterion_id,
            status="complete",
            evidence_found=True,
            feedback="Cumplimiento total.",
        )],
    ))
    ai_audit = client.post(
        f"/api/projects/{project_id}/audit",
        json={"html_content": html, "use_ai": True},
    )
    assert ai_audit.status_code == 200
    ai_item = ai_audit.json()["items"][0]
    assert ai_item["evidence_found"] is False
    assert ai_item["status"] == "partial"

    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    )
    upload = client.post(
        f"/api/projects/{project_id}/evidence",
        files={"file": ("captura.png", BytesIO(png), "image/png")},
        data={"caption": "Ejecución de servicios", "criterion_id": criterion_id},
    )
    assert upload.status_code == 200

    grounded_audit = client.post(
        f"/api/projects/{project_id}/audit",
        json={"html_content": html, "use_ai": True},
    )
    assert grounded_audit.status_code == 200
    grounded_item = grounded_audit.json()["items"][0]
    assert grounded_item["evidence_found"] is True
    assert "Evidencia gráfica vinculada" in grounded_item["feedback"]


def test_sources_with_the_same_filename_keep_distinct_originals() -> None:
    create_resp = client.post("/api/projects/quick-create", json={"title": "Fuentes homónimas"})
    assert create_resp.status_code == 200
    project_id = create_resp.json()["project_id"]

    first = client.post(
        f"/api/projects/{project_id}/sources",
        files={"files": ("material.txt", BytesIO(b"contenido original uno"), "text/plain")},
    )
    second = client.post(
        f"/api/projects/{project_id}/sources",
        files={"files": ("material.txt", BytesIO(b"contenido original dos"), "text/plain")},
    )
    assert first.status_code == 200
    assert second.status_code == 200
    first_source = first.json()["sources"][0]
    second_source = second.json()["sources"][0]
    assert first_source["path"] != second_source["path"]
    assert Path(first_source["path"]).read_bytes() == b"contenido original uno"
    assert Path(second_source["path"]).read_bytes() == b"contenido original dos"

    project = client.get(f"/api/projects/{project_id}").json()
    assert len(project["sources"]) == 2
    assert {source["filename"] for source in project["sources"]} == {"material.txt"}


def test_agent_context_artifacts_and_execution_trace_are_project_scoped() -> None:
    create_resp = client.post("/api/projects/quick-create", json={
        "title": "Proyecto operado por agente",
        "subject": "Base de datos",
        "deliverables": ["Script SQL", "Informe Word"],
        "criteria": [{
            "criterion": "Consultas",
            "indicator": "Construye una consulta SQL verificable",
            "points": 20,
        }],
    })
    project_id = create_resp.json()["project_id"]
    criterion_id = create_resp.json()["criteria"][0]["id"]

    initial_context = client.get(f"/api/agent/projects/{project_id}/context")
    assert initial_context.status_code == 200
    context_data = initial_context.json()
    assert context_data["schema_version"] == "1.0"
    assert context_data["project"]["deliverables"] == ["Script SQL", "Informe Word"]
    assert [item["id"] for item in context_data["pending_criteria"]] == [criterion_id]
    assert context_data["constraints"]["completion_is_audit_controlled"] is True

    artifact_response = client.post(
        f"/api/agent/projects/{project_id}/artifacts",
        files={"file": ("consulta.sql", BytesIO(b"SELECT 1;"), "application/sql")},
        data={
            "kind": "code",
            "criterion_id": criterion_id,
            "source": "codex",
            "description": "Consulta ejecutada en entorno local",
        },
    )
    assert artifact_response.status_code == 200
    artifact = artifact_response.json()
    assert artifact["validation_status"] == "unverified"
    assert artifact["criterion_id"] == criterion_id
    assert client.get(artifact["url"]).content == b"SELECT 1;"

    execution_response = client.post(
        f"/api/agent/projects/{project_id}/executions",
        json={
            "agent": "codex",
            "action": "Ejecutar consulta de prueba",
            "status": "succeeded",
            "criterion_id": criterion_id,
            "command": "sqlite3 test.db < consulta.sql",
            "exit_code": 0,
            "stdout_excerpt": "1",
            "artifact_ids": [artifact["id"]],
        },
    )
    assert execution_response.status_code == 200
    assert execution_response.json()["artifact_ids"] == [artifact["id"]]

    final_context = client.get(f"/api/agent/projects/{project_id}/context").json()
    assert [item["id"] for item in final_context["artifacts"]] == [artifact["id"]]
    assert final_context["recent_executions"][0]["id"] == execution_response.json()["id"]
    assert final_context["project"]["criteria"][0]["status"] == "pending"

    inconsistent_success = client.post(
        f"/api/agent/projects/{project_id}/executions",
        json={
            "agent": "codex",
            "action": "Ejecución fallida",
            "status": "succeeded",
            "exit_code": 1,
        },
    )
    assert inconsistent_success.status_code == 400

    other_project = client.post("/api/projects/quick-create", json={"title": "Otro proyecto"}).json()["project_id"]
    cross_project = client.post(
        f"/api/agent/projects/{other_project}/executions",
        json={
            "agent": "codex",
            "action": "Intento cruzado",
            "status": "partial",
            "artifact_ids": [artifact["id"]],
        },
    )
    assert cross_project.status_code == 400


def test_offline_project_import_extracts_source_without_inventing_criteria(monkeypatch) -> None:
    def fail_if_ai_is_called(*args, **kwargs):
        raise AssertionError(f"La importación local no debe invocar IA: {args!r} {kwargs!r}")

    monkeypatch.setattr(main.ai, "analyze_rubric", fail_if_ai_is_called)
    response = client.post(
        "/api/projects/import",
        files={
            "file": (
                "pauta.txt",
                BytesIO(b"Entregar un informe Word y un script SQL con evidencia verificable."),
                "text/plain",
            )
        },
        data={
            "title": "Evaluación importada localmente",
            "subject": "Consulta de Datos",
            "student": "Nicolás",
            "career": "Informática",
        },
    )
    assert response.status_code == 200
    data = response.json()
    assert data["ai_used"] is False
    assert data["requires_manual_rubric_review"] is True
    assert data["criteria"] == []

    project = client.get(f"/api/projects/{data['project_id']}").json()
    assert project["title"] == "Evaluación importada localmente"
    assert project["criteria"] == []
    assert len(project["sources"]) == 1
    assert project["sources"][0]["filename"] == "pauta.txt"
    assert Path(project["sources"][0]["path"]).read_text(encoding="utf-8").startswith("Entregar un informe")


def test_regenerate_section_endpoint(monkeypatch) -> None:
    create_resp = client.post("/api/projects/quick-create", json={
        "title": "Proyecto Prueba Regeneración",
        "subject": "Redes",
        "student": "Estudiante",
        "pauta_text": "Configurar router y switch",
        "deliverables": ["Informe"],
        "criteria": []
    })
    assert create_resp.status_code == 200
    pid = create_resp.json()["project_id"]

    client.put(f"/api/projects/{pid}/document", json={
        "html_content": "<h2>Introducción</h2><p>Texto viejo de intro.</p><h2>Desarrollo</h2><p>Texto desarrollo.</p>"
    })

    monkeypatch.setattr(main.ai, "text", lambda prompt: "<h2>Introducción</h2><p>Texto generado de nuevo con IA.</p>")

    regen_resp = client.post(f"/api/projects/{pid}/regenerate-section", json={
        "heading": "Introducción",
        "instruction": "Enfocar en topología de red"
    })
    assert regen_resp.status_code == 200
    data = regen_resp.json()
    assert data["status"] == "regenerated"
    assert "Texto generado de nuevo con IA." in data["html_content"]
    assert "Texto desarrollo." in data["html_content"]
    assert data["version"] is not None


def test_regenerate_section_preserves_unsaved_client_html(monkeypatch) -> None:
    from main import store
    create_resp = client.post("/api/projects/quick-create", json={
        "title": "Proyecto Preservación Edición",
        "subject": "Sistemas",
        "student": "Estudiante",
        "pauta_text": "Pauta de prueba",
    })
    pid = create_resp.json()["project_id"]

    # Initial persisted content in SQLite
    client.put(f"/api/projects/{pid}/document", json={
        "html_content": "<h2>Introducción</h2><p>Intro vieja en BD.</p><h2>Desarrollo</h2><p>Desarrollo original en BD.</p>"
    })

    monkeypatch.setattr(main.ai, "text", lambda prompt: "<h2>Introducción</h2><p>Intro regenerada por IA.</p>")

    # Unsaved client-side edit in "Desarrollo"
    unsaved_client_html = (
        "<h2>Introducción</h2><p>Intro vieja en BD.</p>"
        "<h2>Desarrollo</h2><p>Mi edición reciente escrita en Tiptap sin guardar aún en SQLite.</p>"
    )

    regen_resp = client.post(f"/api/projects/{pid}/regenerate-section", json={
        "heading": "Introducción",
        "instruction": "Actualizar",
        "current_html_content": unsaved_client_html,
    })
    assert regen_resp.status_code == 200
    data = regen_resp.json()

    # Verify both the regenerated section AND the unsaved edit are present!
    assert "Intro regenerada por IA." in data["html_content"]
    assert "Mi edición reciente escrita en Tiptap sin guardar aún en SQLite." in data["html_content"]

    # Verify safety backup snapshot was created before applying regen
    versions = store.list_versions(pid)
    assert any("before_regen" in v["reason"] for v in versions)


def test_regenerate_section_empty_heading_rejected() -> None:
    create_resp = client.post("/api/projects/quick-create", json={"title": "Test Empty Heading"})
    pid = create_resp.json()["project_id"]

    resp = client.post(f"/api/projects/{pid}/regenerate-section", json={
        "heading": "   ",
    })
    assert resp.status_code == 400


def test_evidence_watermark_upload() -> None:
    buf = BytesIO()
    img = Image.new("RGBA", (200, 100), color=(30, 41, 59, 255))
    img.save(buf, format="PNG")
    buf.seek(0)

    create_resp = client.post("/api/projects/quick-create", json={
        "title": "Proyecto con Evidencia Watermark",
        "subject": "Sistemas Operativos",
        "student": "Estudiante Prueba",
        "criteria": []
    })
    assert create_resp.status_code == 200
    pid = create_resp.json()["project_id"]

    upload_resp = client.post(
        f"/api/projects/{pid}/evidence",
        files={"file": ("terminal.png", buf, "image/png")},
        data={
            "caption": "Terminal bash",
            "watermark": "true",
            "watermark_text": "Estudiante Prueba - DocStudio - Bimestre 1",
        }
    )
    assert upload_resp.status_code == 200
    ev_data = upload_resp.json()
    assert ev_data["caption"] == "Terminal bash"
    assert ev_data["url"].startswith("/api/evidence/")

    content_resp = client.get(ev_data["url"])
    assert content_resp.status_code == 200
    watermarked_img = Image.open(BytesIO(content_resp.content))
    assert watermarked_img.size == (200, 100)


def test_package_submission_zip() -> None:
    create_resp = client.post("/api/projects/quick-create", json={
        "title": "Proyecto Empaquetado Final",
        "subject": "Bases de Datos",
        "student": "Estudiante Prueba",
        "deliverables": ["Informe Word", "Script SQL"],
        "criteria": []
    })
    assert create_resp.status_code == 200
    pid = create_resp.json()["project_id"]

    sql_content = b"CREATE TABLE usuarios (id INT PRIMARY KEY, nombre VARCHAR(100));"
    client.post(
        f"/api/projects/{pid}/sources",
        files={"files": ("schema.sql", BytesIO(sql_content), "text/plain")}
    )

    pkg_resp = client.post(
        f"/api/projects/{pid}/package-submission",
        json={"html_content": "<h2>Modelo Relacional</h2><p>Explicación del script SQL.</p>"}
    )
    assert pkg_resp.status_code == 200
    assert pkg_resp.headers["content-type"] == "application/zip"

    zf = zipfile.ZipFile(BytesIO(pkg_resp.content))
    namelist = zf.namelist()
    assert "LEEME_ENTREGA.txt" in namelist
    assert any(name.startswith("Informe/") and name.endswith(".docx") for name in namelist)
    assert any("schema.sql" in name for name in namelist)

    manifest = zf.read("LEEME_ENTREGA.txt").decode("utf-8")
    assert "PAQUETE DE ENTREGA ACADÉMICA" in manifest
    assert "Informe Word" in manifest


def test_chat_does_not_auto_insert_memories_without_approval(monkeypatch) -> None:
    from app.models import ChatResult

    create_resp = client.post("/api/projects/quick-create", json={
        "title": "Proyecto Gobernanza Memoria",
        "subject": "Redes",
        "student": "Estudiante",
        "criteria": []
    })
    pid = create_resp.json()["project_id"]

    monkeypatch.setattr(
        main.ai,
        "chat",
        lambda prompt: ChatResult(
            answer="Recomiendo usar topología estrella.",
            suggested_actions=[],
            memories_to_store=["El usuario prefiere topología estrella en capa 2."]
        )
    )

    chat_resp = client.post(f"/api/projects/{pid}/chat", json={
        "message": "¿Qué topología recomiendas?",
        "document_excerpt": ""
    })
    assert chat_resp.status_code == 200
    assert chat_resp.json()["memories_to_store"] == ["El usuario prefiere topología estrella en capa 2."]

    # Verify that memories table in SQLite is still empty (governance preserved)
    memories_resp = client.get(f"/api/projects/{pid}/memories")
    assert memories_resp.status_code == 200
    assert len(memories_resp.json()) == 0

    # User explicitly approves and writes memory
    add_mem_resp = client.post(f"/api/projects/{pid}/memories", json={
        "kind": "decision",
        "content": "El usuario prefiere topología estrella en capa 2."
    })
    assert add_mem_resp.status_code == 200

    memories_resp_after = client.get(f"/api/projects/{pid}/memories")
    assert len(memories_resp_after.json()) == 1
    assert memories_resp_after.json()[0]["content"] == "El usuario prefiere topología estrella en capa 2."


def test_import_project_offline_without_student_and_career_succeeds() -> None:
    # Test that leaving student and career empty does not throw AttributeError: 'NoneType' object has no attribute 'strip'
    response = client.post(
        "/api/projects/import",
        data={
            "title": "Proyecto Sin Datos Personales",
            "pauta_raw_text": "Instrucciones de la pauta de evaluación.",
        },
    )
    assert response.status_code == 200
    data = response.json()
    assert data["title"] == "Proyecto Sin Datos Personales"
    assert data["student"] != ""
    assert data["ai_used"] is False


def test_update_html_section_does_not_duplicate_heading() -> None:
    create_resp = client.post("/api/projects/quick-create", json={"title": "Test Deduplicate Heading"})
    pid = create_resp.json()["project_id"]

    # Initial HTML with heading
    client.put(f"/api/projects/{pid}/document", json={"html_content": "<h2>Introducción</h2><p>Contenido anterior</p>"})

    # Update with replacement that already starts with <h2>
    patch_resp = client.patch(
        f"/api/projects/{pid}/section",
        json={
            "heading": "Introducción",
            "content_html": "<h2>Introducción</h2><p>Contenido nuevo regenerado</p>",
            "create_snapshot": False,
        },
    )
    assert patch_resp.status_code == 200
    html_res = patch_resp.json()["html_content"]
    assert html_res.lower().count("<h2>") == 1
    assert "Contenido nuevo regenerado" in html_res
    assert "Contenido anterior" not in html_res


def test_update_project_outline_persists_to_db() -> None:
    create_resp = client.post("/api/projects/quick-create", json={"title": "Test Persist Outline"})
    pid = create_resp.json()["project_id"]

    new_outline = [
        {"id": "sec-1", "title": "Requerimiento 1: PL/SQL", "description": "Código y capturas", "type": "development", "needs_code": True, "needs_evidence": True},
        {"id": "sec-2", "title": "Conclusión Técnica", "description": "Resumen", "type": "conclusion", "needs_code": False, "needs_evidence": False},
    ]

    put_resp = client.put(f"/api/projects/{pid}/outline", json={"outline": new_outline})
    assert put_resp.status_code == 200

    # Re-fetch project to ensure SQLite persisted the change
    project = client.get(f"/api/projects/{pid}").json()
    assert len(project["outline"]) == 2
    assert project["outline"][0]["title"] == "Requerimiento 1: PL/SQL"
    assert project["outline"][1]["title"] == "Conclusión Técnica"


def test_search_sources_fallback_without_fts5(tmp_path: Path) -> None:
    create_resp = client.post("/api/projects/quick-create", json={"title": "Test Fallback Search"})
    pid = create_resp.json()["project_id"]

    source_content = b"Configuracion avanzada de servicios en Linux RedHat con MariaDB y VSFTPD."
    client.post(
        f"/api/projects/{pid}/sources",
        files={"files": ("redhat.txt", BytesIO(source_content), "text/plain")},
    )

    original_fts = main.store.fts_enabled
    try:
        main.store.fts_enabled = False
        # Search for non-adjacent tokens in fallback
        results = main.store.search_sources(pid, "Linux MariaDB", limit=5)
        assert len(results) > 0
        assert "MariaDB" in results[0]["content"]
    finally:
        main.store.fts_enabled = original_fts


def test_source_chunks_fts_trigger_cleans_deleted_chunks() -> None:
    if not main.store.fts_enabled:
        return
    create_resp = client.post("/api/projects/quick-create", json={"title": "Test FTS Trigger"})
    pid = create_resp.json()["project_id"]

    client.post(
        f"/api/projects/{pid}/sources",
        files={"files": ("prueba_trigger.txt", BytesIO(b"Contenido especifico para validar trigger de limpieza"), "text/plain")},
    )

    with main.store.connect() as db:
        before_count = db.execute("SELECT COUNT(*) FROM source_chunks_fts WHERE project_id=?", (pid,)).fetchone()[0]
        assert before_count > 0
        # Delete chunk from parent table
        db.execute("DELETE FROM source_chunks WHERE project_id=?", (pid,))
        after_count = db.execute("SELECT COUNT(*) FROM source_chunks_fts WHERE project_id=?", (pid,)).fetchone()[0]
        assert after_count == 0


def test_delete_project_and_source_endpoints() -> None:
    create_resp = client.post("/api/projects/quick-create", json={"title": "Proyecto a Eliminar API"})
    assert create_resp.status_code == 200
    pid = create_resp.json()["project_id"]

    # Upload a source
    src_resp = client.post(
        f"/api/projects/{pid}/sources",
        files={"files": ("temp_doc.txt", BytesIO(b"Documento de prueba para borrar"), "text/plain")},
    )
    assert src_resp.status_code == 200
    src_id = src_resp.json()["sources"][0]["id"]

    # Delete source endpoint
    del_src = client.delete(f"/api/projects/{pid}/sources/{src_id}")
    assert del_src.status_code == 200
    assert del_src.json()["status"] == "deleted"

    # Delete project endpoint
    del_prj = client.delete(f"/api/projects/{pid}")
    assert del_prj.status_code == 200
    assert del_prj.json()["status"] == "deleted"

    # Verify 404 after deletion
    get_prj = client.get(f"/api/projects/{pid}")
    assert get_prj.status_code == 404


def test_agent_export_docx_and_package_submission_endpoints() -> None:
    create_resp = client.post("/api/projects/quick-create", json={
        "title": "Proyecto Test Export Agente",
        "subject": "Redes",
        "student": "Agente Antigravity",
    })
    assert create_resp.status_code == 200
    pid = create_resp.json()["project_id"]

    # 1. Agent export DOCX endpoint
    docx_resp = client.post(f"/api/agent/projects/{pid}/export-docx")
    assert docx_resp.status_code == 200
    docx_data = docx_resp.json()
    assert docx_data["status"] == "ok"
    assert docx_data["project_id"] == pid
    assert Path(docx_data["file_path"]).is_file()
    assert docx_data["size_bytes"] > 0

    # 2. Agent package submission ZIP endpoint
    pkg_resp = client.post(f"/api/agent/projects/{pid}/package-submission")
    assert pkg_resp.status_code == 200
    pkg_data = pkg_resp.json()
    assert pkg_data["status"] == "ok"
    assert pkg_data["project_id"] == pid
    assert Path(pkg_data["zip_path"]).is_file()
    assert pkg_data["size_bytes"] > 0

    # 3. Verify checksums.sha256 in the generated ZIP
    with zipfile.ZipFile(pkg_data["zip_path"], "r") as zf:
        namelist = zf.namelist()
        assert "checksums.sha256" in namelist
        checksums_content = zf.read("checksums.sha256").decode("utf-8")
        assert len(checksums_content.strip().splitlines()) >= 2
        for line in checksums_content.strip().splitlines():
            digest, arcname = line.split("  ", 1)
            assert len(digest) == 64
            assert arcname in namelist


def test_agent_artifact_image_preserved_in_docx() -> None:
    create_resp = client.post("/api/projects/quick-create", json={
        "title": "Proyecto Artefacto Grafico",
        "subject": "Arquitectura",
    })
    pid = create_resp.json()["project_id"]

    # Generate a real 100x100 PNG
    img_buf = BytesIO()
    Image.new("RGB", (100, 100), color="blue").save(img_buf, format="PNG")
    img_bytes = img_buf.getvalue()

    # Agent attaches artifact
    art_resp = client.post(
        f"/api/agent/projects/{pid}/artifacts",
        files={"file": ("diagrama_red.png", BytesIO(img_bytes), "image/png")},
        data={"kind": "image", "source": "antigravity", "description": "Diagrama de arquitectura"},
    )
    assert art_resp.status_code == 200
    art_id = art_resp.json()["id"]

    # Put document containing the agent artifact URL
    html_with_artifact = f'<h2>Arquitectura</h2><p><img src="/api/agent/artifacts/{art_id}/content" alt="Diagrama de red" /></p>'
    client.put(f"/api/projects/{pid}/document", json={"html_content": html_with_artifact})

    # Export DOCX
    exp_resp = client.post(f"/api/agent/projects/{pid}/export-docx")
    assert exp_resp.status_code == 200
    docx_file = Path(exp_resp.json()["file_path"])
    assert docx_file.is_file()

    # Inspect generated document: must NOT contain EVIDENCIA PENDIENTE!
    doc = Document(str(docx_file))
    all_text = " ".join(p.text for p in doc.paragraphs)
    assert "EVIDENCIA PENDIENTE" not in all_text
    assert "Diagrama de red" in all_text


def test_export_docx_and_package_with_webp_evidence() -> None:
    create_resp = client.post("/api/projects/quick-create", json={
        "title": "Proyecto Evidencia WebP",
        "subject": "Desarrollo Seguro",
    })
    assert create_resp.status_code == 200
    pid = create_resp.json()["project_id"]

    # Generate a real WebP image in memory
    webp_buf = BytesIO()
    Image.new("RGBA", (120, 120), color=(20, 100, 200, 255)).save(webp_buf, format="WEBP")
    webp_bytes = webp_buf.getvalue()

    # Upload as evidence (watermark=False to keep raw WebP on disk)
    ev_resp = client.post(
        f"/api/projects/{pid}/evidence",
        files={"file": ("captura_terminal.webp", BytesIO(webp_bytes), "image/webp")},
        data={"caption": "Terminal con migraciones", "watermark": "false"},
    )
    assert ev_resp.status_code == 200
    ev_id = ev_resp.json()["id"]

    # Save document referencing the WebP evidence URL
    doc_html = f'<h2>Migraciones</h2><p><img src="/api/evidence/{ev_id}/content" alt="Terminal con migraciones" /></p>'
    client.put(f"/api/projects/{pid}/document", json={"html_content": doc_html})

    # 1. Export DOCX via main export endpoint
    export_resp = client.post("/api/export-docx", json={
        "project_id": pid,
        "title": "Proyecto Evidencia WebP",
        "html_content": doc_html,
    })
    assert export_resp.status_code == 200
    assert export_resp.content.startswith(b"PK")
    doc = Document(BytesIO(export_resp.content))
    all_text = " ".join(p.text for p in doc.paragraphs)
    assert "EVIDENCIA PENDIENTE" not in all_text
    assert "Terminal con migraciones" in all_text
    assert len(doc.inline_shapes) >= 1

    # 2. Package submission ZIP
    pkg_resp = client.post(f"/api/projects/{pid}/package-submission", json={"html_content": doc_html})
    assert pkg_resp.status_code == 200
    assert pkg_resp.content.startswith(b"PK")
    with zipfile.ZipFile(BytesIO(pkg_resp.content)) as zf:
        namelist = zf.namelist()
        assert "LEEME_ENTREGA.txt" in namelist
        assert "checksums.sha256" in namelist
        docx_members = [m for m in namelist if m.startswith("Informe/") and m.endswith(".docx")]
        assert len(docx_members) == 1
        # Also verify that the evidence file itself is packaged in Evidencias/
        ev_members = [m for m in namelist if m.startswith("Evidencias/")]
        assert len(ev_members) == 1


def test_put_and_post_project_criteria() -> None:
    # 1. Create project without criteria (simulating offline import)
    create_resp = client.post("/api/projects/quick-create", json={
        "title": "Proyecto Offline Sin Criterios Iniciales",
        "subject": "Sistemas Operativos",
        "student": "Estudiante DS01",
        "criteria": [],
    })
    assert create_resp.status_code == 200
    pid = create_resp.json()["project_id"]

    # Initial get: criteria is empty
    proj = client.get(f"/api/projects/{pid}").json()
    assert len(proj["criteria"]) == 0

    # 2. PUT criteria to define initial rubric
    initial_criteria = [
        {"criterion": "Configuración SSH", "indicator": "Permitir acceso por llave pública", "points": 10.0, "requires_evidence": True},
        {"criterion": "Cortafuegos", "indicator": "Políticas DROP en iptables", "points": 15.0, "requires_evidence": False},
    ]
    put_resp = client.put(f"/api/projects/{pid}/criteria", json={"criteria": initial_criteria})
    assert put_resp.status_code == 200
    put_data = put_resp.json()
    assert put_data["status"] == "saved"
    assert len(put_data["criteria"]) == 2
    c1_id = put_data["criteria"][0]["id"]
    c2_id = put_data["criteria"][1]["id"]
    assert c1_id.startswith("crit_")
    assert c2_id.startswith("crit_")

    # 3. Associate an evidence file to c1
    ev_resp = client.post(
        f"/api/projects/{pid}/evidence",
        files={"file": ("ssh_config.png", BytesIO(b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15c4\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"), "image/png")},
        data={"caption": "Captura SSH", "criterion_id": c1_id},
    )
    assert ev_resp.status_code == 200
    ev_id = ev_resp.json()["id"]

    # 4. Replace criteria: preserve c1 (updating indicator), drop c2, add c3
    updated_criteria = [
        {"id": c1_id, "criterion": "Configuración SSH Avanzada", "indicator": "Permitir acceso exclusivo por llave ED25519", "points": 12.0, "requires_evidence": True},
        {"criterion": "Auditoría de Logs", "indicator": "Monitoreo con journalctl", "points": 8.0, "requires_evidence": False},
    ]
    put2_resp = client.put(f"/api/projects/{pid}/criteria", json={"criteria": updated_criteria})
    assert put2_resp.status_code == 200
    put2_data = put2_resp.json()
    assert len(put2_data["criteria"]) == 2
    assert put2_data["criteria"][0]["id"] == c1_id
    assert put2_data["criteria"][0]["points"] == 12.0
    c3_id = put2_data["criteria"][1]["id"]
    assert c3_id != c2_id

    # Verify evidence still points to c1_id
    ev_after = main.store.get_evidence(ev_id)
    assert ev_after is not None
    assert ev_after["criterion_id"] == c1_id

    # 5. POST to incrementally add a 3rd criterion
    post_resp = client.post(f"/api/projects/{pid}/criteria", json={
        "criteria": [{"criterion": "Backup Automático", "indicator": "Script bash con cron", "points": 10.0}]
    })
    assert post_resp.status_code == 200
    assert post_resp.json()["status"] == "added"

    proj_final = client.get(f"/api/projects/{pid}").json()
    assert len(proj_final["criteria"]) == 3
    assert proj_final["criteria"][2]["position"] == 2
    assert proj_final["criteria"][2]["criterion"] == "Backup Automático"

    # 6. Run audit and verify newly added criteria are audited
    audit_resp = client.post(f"/api/projects/{pid}/audit", json={
        "html_content": "<p>Permitir acceso exclusivo por llave ED25519 con backup automático.</p>",
        "use_ai": False,
    })
    assert audit_resp.status_code == 200
    audit_data = audit_resp.json()
    assert len(audit_data["items"]) == 3


def test_upload_evidence_with_empty_criterion_id_succeeds_and_stores_null() -> None:
    create_resp = client.post("/api/projects/quick-create", json={"title": "Test Evidencia Sin Criterio"})
    assert create_resp.status_code == 200
    pid = create_resp.json()["project_id"]

    png_bytes = (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00"
        b"\x1f\x15c4\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
    )

    # Empty string criterion_id should not crash with 500 SQLite FK IntegrityError
    upload_resp = client.post(
        f"/api/projects/{pid}/evidence",
        files={"file": ("captura_vacia.png", BytesIO(png_bytes), "image/png")},
        data={"caption": "Captura libre", "criterion_id": ""},
    )
    assert upload_resp.status_code == 200
    ev_data = upload_resp.json()
    assert ev_data["criterion_id"] is None
    assert ev_data["url"].startswith("/api/evidence/")


def test_upload_evidence_with_invalid_or_cross_project_criterion_id_returns_400_and_cleans_file() -> None:
    from app.settings import settings

    create_resp = client.post("/api/projects/quick-create", json={"title": "Test Evidencia Invalida"})
    assert create_resp.status_code == 200
    pid = create_resp.json()["project_id"]

    other_resp = client.post("/api/projects/quick-create", json={
        "title": "Otro Proyecto",
        "criteria": [{"criterion": "Criterio Ajeno", "indicator": "Indicador Ajeno"}],
    })
    assert other_resp.status_code == 200
    other_crit_id = other_resp.json()["criteria"][0]["id"]

    png_bytes = (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00"
        b"\x1f\x15c4\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
    )

    ev_dir = settings.projects_dir / pid / "evidence"

    # 1. Invalid criterion ID
    bad_resp = client.post(
        f"/api/projects/{pid}/evidence",
        files={"file": ("captura_bad.png", BytesIO(png_bytes), "image/png")},
        data={"caption": "Captura mala", "criterion_id": "crit_inexistente_123"},
    )
    assert bad_resp.status_code == 400
    assert "El criterio no pertenece al proyecto" in bad_resp.json()["detail"]
    # Check no orphan files left in evidence dir
    if ev_dir.exists():
        assert len(list(ev_dir.iterdir())) == 0

    # 2. Cross-project criterion ID
    cross_resp = client.post(
        f"/api/projects/{pid}/evidence",
        files={"file": ("captura_cross.png", BytesIO(png_bytes), "image/png")},
        data={"caption": "Captura cruzada", "criterion_id": other_crit_id},
    )
    assert cross_resp.status_code == 400
    assert "El criterio no pertenece al proyecto" in cross_resp.json()["detail"]
    if ev_dir.exists():
        assert len(list(ev_dir.iterdir())) == 0

