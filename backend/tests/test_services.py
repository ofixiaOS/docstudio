from pathlib import Path
from tempfile import TemporaryDirectory
from fastapi import UploadFile
import io

from app.html_utils import (
    clean_model_output,
    parse_html_to_sections,
    resolve_image_path,
    update_html_section,
)
from app.prompts import default_outline, section_prompt
from app.uploads import save_upload, save_content_addressed_upload
from app.audit_service import audit_score, offline_audit, valid_evidence_by_criterion
from app.models import AuditItem
from docx import Document
from docx_exporter import create_docx_document
from PIL import Image


def test_clean_model_output() -> None:
    assert clean_model_output("```html\n<h2>Titulo</h2>\n```") == "<h2>Titulo</h2>"
    assert clean_model_output("```json\n{\"ok\": true}\n```") == "{\"ok\": true}"
    assert clean_model_output("<p>Normal</p>") == "<p>Normal</p>"


def test_update_html_section_append_and_replace() -> None:
    initial = "<h2>Introducción</h2><p>Texto inicial</p>"
    updated = update_html_section(initial, "Introducción", "<p>Texto reemplazado</p>")
    assert "Texto reemplazado" in updated
    assert "Texto inicial" not in updated

    appended = update_html_section(initial, "Conclusión", "<p>Texto nuevo</p>")
    assert "Introducción" in appended
    assert "Conclusión" in appended
    assert "Texto nuevo" in appended


def test_parse_html_to_sections() -> None:
    html = "<h1>Titulo</h1><p>Parrafo</p><pre><code>codigo</code></pre>"
    sections = parse_html_to_sections(html)
    assert len(sections) == 3
    assert sections[0]["type"] == "h1"
    assert sections[1]["type"] == "paragraph"
    assert sections[2]["type"] == "code"


def test_resolve_image_path_custom_getter() -> None:
    fake_evidence = {"stored_path": "C:/fake/path.png"}
    def fake_getter(eid: str):
        return fake_evidence if eid == "ev123" else None

    assert resolve_image_path("/api/evidence/ev123/content", get_evidence=fake_getter) == "C:/fake/path.png"
    assert resolve_image_path("/api/evidence/unknown/content", get_evidence=fake_getter) is None
    assert resolve_image_path(None, get_evidence=fake_getter) is None


def test_prompts_default_outline_and_section_prompt() -> None:
    outline = default_outline()
    assert len(outline) == 4
    assert outline[0]["title"] == "Introducción"

    card = {"title": "Desarrollo", "description": "Detalle técnico"}
    class DummyReq:
        title = "Mi Trabajo"
        subject = "Programación"

    prompt = section_prompt(DummyReq(), card, "criterio 1", "fuente 1")
    assert "Mi Trabajo" in prompt
    assert "Programación" in prompt
    assert "Desarrollo" in prompt


def test_audit_score_and_offline_audit() -> None:
    project = {
        "criteria": [
            {"id": "c1", "indicator": "Desarrollar diagramas de arquitectura", "points": 10, "requires_evidence": False},
            {"id": "c2", "indicator": "Implementar backend con pruebas unitarias", "points": 20, "requires_evidence": True},
        ],
        "evidence": [],
    }
    html = "<h2>Arquitectura</h2><p>Diagramas de arquitectura detallados.</p>"
    result = offline_audit(project, html)
    assert len(result.items) == 2
    assert result.items[0].criterion_id == "c1"
    assert result.items[0].status == "complete"
    assert result.items[1].criterion_id == "c2"
    assert result.items[1].status == "missing"
    assert result.estimated_score > 0


def test_resolve_image_path_with_agent_artifact() -> None:
    fake_artifact = {"id": "art_123", "stored_path": "C:/fake/diagram.png"}
    def fake_artifact_getter(aid: str):
        return fake_artifact if aid == "art_123" else None

    assert resolve_image_path("/api/agent/artifacts/art_123/content", get_artifact=fake_artifact_getter) == "C:/fake/diagram.png"
    assert resolve_image_path("/api/agent/artifacts/missing/content", get_artifact=fake_artifact_getter) is None


def test_valid_evidence_recognizes_artifacts_and_executions(tmp_path: Path) -> None:
    art_file = tmp_path / "diagram.png"
    art_file.write_text("fake diagram", encoding="utf-8")
    project = {
        "criteria": [
            {"id": "c1", "indicator": "Diagrama de BD", "points": 10, "requires_evidence": True},
            {"id": "c2", "indicator": "Pruebas unitarias", "points": 10, "requires_evidence": True},
            {"id": "c3", "indicator": "Artefacto corrupto", "points": 10, "requires_evidence": True},
        ],
        "evidence": [],
        "artifacts": [
            {"id": "a1", "criterion_id": "c1", "stored_path": str(art_file), "validation_status": "unverified", "filename": "diagram.png"},
            {"id": "a2", "criterion_id": "c3", "stored_path": str(tmp_path / "missing.png"), "validation_status": "unverified", "filename": "missing.png"},
        ],
        "executions": [
            {"id": "e1", "criterion_id": "c2", "status": "succeeded", "agent": "antigravity", "action": "run_pytest"},
            {"id": "e2", "criterion_id": "c3", "status": "failed", "agent": "antigravity", "action": "fail_test"},
        ],
    }
    evidence_map = valid_evidence_by_criterion(project)
    assert "c1" in evidence_map
    assert evidence_map["c1"][0]["type"] == "artifact"
    assert "c2" in evidence_map
    assert evidence_map["c2"][0]["type"] == "execution"
    # Fail-closed check: c3 only had a missing file and a failed execution, so it must NOT be satisfied!
    assert "c3" not in evidence_map


def test_docx_add_image_transcodes_webp(tmp_path: Path) -> None:
    webp_file = tmp_path / "screenshot.webp"
    Image.new("RGB", (60, 60), color="purple").save(webp_file, format="WEBP")
    out_docx = tmp_path / "output_webp.docx"
    doc_data = {
        "title": "Informe con WebP",
        "sections": [
            {"type": "h1", "content": "Evidencia Gráfica"},
            {"type": "image", "image_path": str(webp_file), "image_alt": "Captura de pruebas"},
        ],
    }
    create_docx_document(doc_data, str(out_docx))
    assert out_docx.is_file()
    doc = Document(str(out_docx))
    all_text = " ".join(p.text for p in doc.paragraphs)
    assert "EVIDENCIA PENDIENTE" not in all_text
    assert "Captura de pruebas" in all_text
    assert len(doc.inline_shapes) >= 1


def test_docx_add_image_handles_corrupt_file_gracefully(tmp_path: Path) -> None:
    zero_byte_file = tmp_path / "corrupt.png"
    zero_byte_file.write_bytes(b"")
    out_docx = tmp_path / "output_corrupt.docx"
    doc_data = {
        "title": "Informe con Archivo Dañado",
        "sections": [
            {"type": "image", "image_path": str(zero_byte_file), "image_alt": "Captura corrupta"},
            {"type": "image", "image_path": str(tmp_path / "nonexistent.png"), "image_alt": "Captura borrada"},
        ],
    }
    # Must NOT raise any exception; must degrade cleanly to EVIDENCIA PENDIENTE
    create_docx_document(doc_data, str(out_docx))
    assert out_docx.is_file()
    doc = Document(str(out_docx))
    all_text = " ".join(p.text for p in doc.paragraphs)
    assert "EVIDENCIA PENDIENTE: Captura corrupta" in all_text
    assert "EVIDENCIA PENDIENTE: Captura borrada" in all_text


def test_update_html_section_empty_heading_and_numeric_collision() -> None:
    html = "<h1>Documento Principal</h1><p>Texto raíz</p><h2>1. Introducción</h2><p>Intro original</p>"

    # 1. Empty or whitespace heading must NOT overwrite anything
    assert update_html_section(html, "", "<p>Ignorar</p>") == html
    assert update_html_section(html, "   ", "<p>Ignorar</p>") == html

    # 2. Short numeric heading "1" must NOT collide with "1. Introducción"
    res_num = update_html_section(html, "1", "<p>Sección numérica</p>")
    assert "Documento Principal" in res_num
    assert "Intro original" in res_num
    assert "<h2>1</h2>" in res_num
    assert "<p>Sección numérica</p>" in res_num

    # 3. Numbered title matching: "1. Introducción" matches "1. Introducción", and "Introducción" matches "1. Introducción"
    res_norm = update_html_section(html, "Introducción", "<p>Intro actualizada</p>")
    assert "Intro actualizada" in res_norm
    assert "Intro original" not in res_norm
    assert "Documento Principal" in res_norm


def test_parse_html_to_sections_preserves_h4_to_h6_and_container_text() -> None:
    html = (
        "<div>Texto directo en div contenedor</div>"
        "<section>Texto en section sin bloque</section>"
        "<h4>1.1.1 Subsección Nivel 4</h4>"
        "<h5>1.1.1.1 Nivel 5</h5>"
        "<h6>1.1.1.1.1 Nivel 6</h6>"
    )
    sections = parse_html_to_sections(html)
    assert len(sections) == 5
    assert sections[0]["type"] == "paragraph"
    assert sections[0]["content"] == "Texto directo en div contenedor"
    assert sections[1]["type"] == "paragraph"
    assert sections[1]["content"] == "Texto en section sin bloque"
    assert sections[2]["type"] == "h4"
    assert sections[2]["content"] == "1.1.1 Subsección Nivel 4"
    assert sections[3]["type"] == "h5"
    assert sections[3]["content"] == "1.1.1.1 Nivel 5"
    assert sections[4]["type"] == "h6"
    assert sections[4]["content"] == "1.1.1.1.1 Nivel 6"


def test_update_html_section_respects_heading_hierarchy_with_subsections() -> None:
    html = (
        "<h2>Desarrollo</h2>"
        "<p>Intro vieja</p>"
        "<h3>1.1 Requerimiento Técnico</h3>"
        "<p>Detalle viejo</p>"
        "<h4>1.1.1 Especificación</h4>"
        "<p>Subdetalle viejo</p>"
        "<h2>Conclusión</h2>"
        "<p>Texto intacto de conclusión</p>"
    )
    # Updating <h2>Desarrollo</h2> must replace everything up to the next <h2> (including <h3> and <h4>)
    updated = update_html_section(
        html,
        "Desarrollo",
        "<h2>Desarrollo</h2><p>Nuevo desarrollo integral con arquitectura hexagonal.</p>",
    )
    assert "Nuevo desarrollo integral con arquitectura hexagonal." in updated
    assert "Intro vieja" not in updated
    assert "1.1 Requerimiento Técnico" not in updated
    assert "1.1.1 Especificación" not in updated
    assert "Conclusión" in updated
    assert "Texto intacto de conclusión" in updated
    # Verify no heading duplication
    assert updated.count("<h2>Desarrollo</h2>") == 1


def test_docx_export_renders_h4_to_h6_without_loss(tmp_path: Path) -> None:
    out_docx = tmp_path / "output_hierarchy.docx"
    doc_data = {
        "title": "Informe Técnico con Jerarquía Completa",
        "sections": [
            {"type": "h1", "content": "1. Arquitectura General"},
            {"type": "h2", "content": "1.1 Componentes del Sistema"},
            {"type": "h3", "content": "1.1.1 Servicios de Red"},
            {"type": "h4", "content": "1.1.1.1 Servidor Web Nginx"},
            {"type": "h5", "content": "1.1.1.1.1 Configuración de Worker Processes"},
            {"type": "h6", "content": "1.1.1.1.1.1 Parámetro worker_connections"},
            {"type": "paragraph", "content": "Detalle técnico de configuración final."},
        ],
    }
    create_docx_document(doc_data, str(out_docx))
    assert out_docx.is_file()

    doc = Document(str(out_docx))
    headings_found = {}
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading"):
            headings_found[p.style.name] = p.text

    assert "Heading 1" in headings_found
    assert "Heading 2" in headings_found
    assert "Heading 3" in headings_found
    assert "Heading 4" in headings_found
    assert "Heading 5" in headings_found
    assert "Heading 6" in headings_found
    assert headings_found["Heading 4"] == "1.1.1.1 Servidor Web Nginx"
    assert headings_found["Heading 5"] == "1.1.1.1.1 Configuración de Worker Processes"
    assert headings_found["Heading 6"] == "1.1.1.1.1.1 Parámetro worker_connections"


