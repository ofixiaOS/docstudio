from __future__ import annotations

import hashlib
import re
from pathlib import Path

import pypdf
from docx import Document


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".csv", ".sql", ".py", ".java", ".cs"}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def extract_text(path: Path, max_chars: int = 120_000) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        reader = pypdf.PdfReader(str(path))
        parts = []
        for number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                parts.append(f"[Página {number}]\n{text.strip()}")
            if sum(len(part) for part in parts) >= max_chars:
                break
        result = "\n\n".join(parts)
    elif suffix == ".docx":
        doc = Document(str(path))
        parts = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text.strip())
        for table in doc.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells]
                if any(values):
                    parts.append(" | ".join(values))
        result = "\n".join(parts)
    elif suffix in SUPPORTED_EXTENSIONS:
        result = path.read_text(encoding="utf-8", errors="replace")
    else:
        raise ValueError(f"Formato no soportado: {suffix or 'sin extensión'}")
    return result[:max_chars].strip()


def chunk_text(text: str, target_chars: int = 1800, overlap_chars: int = 240) -> list[str]:
    clean = re.sub(r"[ \t]+", " ", text)
    clean = re.sub(r"\n{3,}", "\n\n", clean).strip()
    if not clean:
        return []
    paragraphs = [part.strip() for part in clean.split("\n\n") if part.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        candidate = f"{current}\n\n{paragraph}".strip()
        if current and len(candidate) > target_chars:
            chunks.append(current)
            tail = current[-overlap_chars:] if overlap_chars else ""
            current = f"{tail}\n\n{paragraph}".strip()
        else:
            current = candidate
    if current:
        chunks.append(current)
    return chunks


def classify_document(path: Path) -> str:
    name = path.name.lower()
    if any(token in name for token in ("material", "me_", "comp_", "tutorial")):
        return "study_material"
    if any(token in name for token in ("evaluaci", "instructivo", "ta_", "examen", "foro")):
        return "rubric"
    if any(token in name for token in ("nicolas", "nicolás", "desarrollo")):
        return "submission"
    return "other"


def list_library_files(root: Path, limit: int = 500) -> list[dict]:
    if not root.exists():
        return []
    files: list[dict] = []
    for path in root.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue
        if "gamma-iplacex" in path.parts:
            continue
        stat = path.stat()
        files.append(
            {
                "path": str(path),
                "relative_path": str(path.relative_to(root)),
                "name": path.name,
                "extension": path.suffix.lower(),
                "kind": classify_document(path),
                "size": stat.st_size,
                "modified_at": stat.st_mtime,
            }
        )
    files.sort(key=lambda item: item["modified_at"], reverse=True)
    return files[:limit]
