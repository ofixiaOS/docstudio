"""Deterministic Word exporter for technical and academic documents."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path

from PIL import Image
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor


BLACK = RGBColor(0, 0, 0)
TEXT = RGBColor(33, 37, 41)
MUTED = RGBColor(70, 80, 95)
BLUE = RGBColor(16, 55, 120)
LIGHT_BORDER = "D9D9D9"


def set_cell_background(cell, fill_hex: str) -> None:
    cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))


def set_cell_margins(cell, top=120, bottom=120, left=180, right=180) -> None:
    cell._tc.get_or_add_tcPr().append(
        parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/>'
            f'<w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/>'
            f'</w:tcMar>'
        )
    )


def set_table_borders(table) -> None:
    table._tbl.tblPr.append(
        parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            + "".join(
                f'<w:{edge} w:val="single" w:sz="4" w:space="0" w:color="{LIGHT_BORDER}"/>'
                for edge in ("top", "left", "bottom", "right", "insideH", "insideV")
            )
            + '</w:tblBorders>'
        )
    )


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15
    title = doc.styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = BLACK
    heading_configs = (
        (1, 16, 14, 5),
        (2, 13, 10, 5),
        (3, 11.5, 10, 4),
        (4, 11, 8, 3),
        (5, 10.5, 6, 3),
        (6, 10, 6, 2),
    )
    for level, size, before, after in heading_configs:
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_html_runs(paragraph, content_html: str | None, fallback: str) -> None:
    if not content_html:
        paragraph.add_run(fallback)
        return
    soup = BeautifulSoup(content_html, "html.parser")

    def visit(node, *, bold=False, italic=False, code=False) -> None:
        if isinstance(node, NavigableString):
            if not str(node):
                return
            run = paragraph.add_run(str(node))
            run.bold = bold
            run.italic = italic
            run.font.name = "Consolas" if code else "Arial"
            run.font.size = Pt(9.5 if code else 11)
            run.font.color.rgb = TEXT
            return
        if not isinstance(node, Tag):
            return
        if node.name == "br":
            paragraph.add_run().add_break()
            return
        next_bold = bold or node.name in {"strong", "b"}
        next_italic = italic or node.name in {"em", "i"}
        next_code = code or node.name == "code"
        for child in node.children:
            visit(child, bold=next_bold, italic=next_italic, code=next_code)

    root = soup.body or soup
    for child in root.children:
        visit(child)


def add_cover(doc: Document, data: dict) -> None:
    institution_text = (data.get("institution") or "").strip()
    career_text = (data.get("career") or "").strip()
    student_text = (data.get("student") or "Autor").strip()
    title_text = (data.get("title") or "Documento Técnico").strip()
    subject_text = (data.get("subject") or "General").strip()
    date_text = data.get("date") or datetime.now().strftime("%d-%m-%Y")

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(36)
    
    if institution_text:
        institution = doc.add_paragraph()
        run = institution.add_run(institution_text.upper())
        run.font.name = "Arial"
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = MUTED

    if career_text:
        career = doc.add_paragraph(career_text.upper())
        career.runs[0].font.name = "Arial"
        career.runs[0].font.size = Pt(10)
        career.runs[0].font.color.rgb = MUTED
        career.paragraph_format.space_after = Pt(40)
    else:
        spacer2 = doc.add_paragraph()
        spacer2.paragraph_format.space_after = Pt(28)

    title = doc.add_paragraph(style="Title")
    title.add_run(title_text)
    title.paragraph_format.space_after = Pt(14)
    
    if subject_text and subject_text.lower() != "general":
        subject = doc.add_paragraph()
        subject_run = subject.add_run(f"Área / Asignatura: {subject_text}")
        subject_run.font.name = "Arial"
        subject_run.font.size = Pt(13)
        subject_run.font.bold = True
        subject_run.font.color.rgb = BLACK
        subject.paragraph_format.space_after = Pt(72)
    else:
        spacer3 = doc.add_paragraph()
        spacer3.paragraph_format.space_after = Pt(60)

    fields = [
        ("Autor", student_text),
    ]
    if subject_text and subject_text.lower() != "general":
        fields.append(("Tema / Área", subject_text))
    fields.append(("Fecha", date_text))

    table = doc.add_table(rows=len(fields), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_borders(table)

    for row, (label, value) in zip(table.rows, fields):
        row.cells[0].width = Inches(1.35)
        row.cells[1].width = Inches(5.0)
        for cell in row.cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_background(row.cells[0], "EAF0F7")
        label_run = row.cells[0].paragraphs[0].add_run(label)
        label_run.bold = True
        label_run.font.name = "Arial"
        label_run.font.size = Pt(10.5)
        value_run = row.cells[1].paragraphs[0].add_run(str(value))
        value_run.font.name = "Arial"
        value_run.font.size = Pt(10.5)
    doc.add_page_break()


def add_code_block(doc: Document, content: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    cell = table.cell(0, 0)
    cell.width = Inches(6.4)
    set_cell_background(cell, "F4F5F7")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(content.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = TEXT
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def add_table_block(doc: Document, rows_data: list[list[str]]) -> None:
    if not rows_data or not any(rows_data):
        return
    cols = max(len(row) for row in rows_data)
    table = doc.add_table(rows=len(rows_data), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    for row_idx, row in enumerate(rows_data):
        is_header = (row_idx == 0)
        for col_idx in range(cols):
            val = row[col_idx] if col_idx < len(row) else ""
            cell = table.cell(row_idx, col_idx)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            if is_header:
                set_cell_background(cell, "EAF0F7")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(val)
            run.font.name = "Arial"
            run.font.size = Pt(9.5)
            if is_header:
                run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_image(doc: Document, path_value: str | None, alt: str | None) -> None:
    if not path_value or not Path(path_value).is_file():
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(f"EVIDENCIA PENDIENTE: {alt or 'Inserte aquí la captura correspondiente.'}")
        run.bold = True
        return

    # Pillow inspection and transparent conversion for WebP or non-native Word formats
    image_source = None
    try:
        with Image.open(path_value) as img:
            img_format = (img.format or "").upper()
            # python-docx supports PNG, JPEG, GIF, TIFF, BMP natively.
            # If WebP or other non-native formats, transcode in-memory to PNG.
            if img_format in {"WEBP", "ICO"} or img_format not in {"PNG", "JPEG", "JPG", "GIF", "TIFF", "BMP"}:
                buf = BytesIO()
                if img.mode in ("RGBA", "LA", "P"):
                    img.save(buf, format="PNG")
                else:
                    img.convert("RGB").save(buf, format="PNG")
                buf.seek(0)
                image_source = buf
            else:
                image_source = path_value
    except Exception:
        image_source = None

    if image_source is None:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(f"EVIDENCIA PENDIENTE: {alt or 'Inserte aquí la captura correspondiente.'}")
        run.bold = True
        return

    try:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(image_source, width=Inches(6.2))
    except Exception:
        paragraph.text = ""
        run = paragraph.add_run(f"EVIDENCIA PENDIENTE: {alt or 'Inserte aquí la captura correspondiente.'}")
        run.bold = True
        return

    if alt:
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_text = alt.strip()
        if not caption_text.lower().startswith("figura"):
            caption_text = f"Figura: {caption_text}"
        caption_run = caption.add_run(caption_text)
        caption_run.italic = True
        caption_run.font.name = "Arial"
        caption_run.font.size = Pt(9)


def create_docx_document(doc_data: dict, output_path: str) -> str:
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.different_first_page_header_footer = True
    add_cover(doc, doc_data)

    for item in doc_data.get("sections", []):
        section_type = item.get("type", "paragraph")
        content = item.get("content", "")
        content_html = item.get("content_html")
        if section_type in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            level = int(section_type[-1])
            paragraph = doc.add_paragraph(style=f"Heading {level}")
            paragraph.add_run(content)
        elif section_type == "code":
            add_code_block(doc, content)
        elif section_type == "table":
            add_table_block(doc, item.get("table_rows") or [])
        elif section_type == "image":
            add_image(doc, item.get("image_path"), item.get("image_alt"))
        elif section_type == "callout":
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(content)
            run.italic = True
            if content.upper().startswith("EVIDENCIA PENDIENTE"):
                run.bold = True
        elif section_type in {"list_item", "ordered_item"}:
            style = "List Number" if section_type == "ordered_item" else "List Bullet"
            paragraph = doc.add_paragraph(style=style)
            add_html_runs(paragraph, content_html, content)
        else:
            paragraph = doc.add_paragraph()
            add_html_runs(paragraph, content_html, content)

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return str(output)


# Export alias
create_studio_document = create_docx_document
